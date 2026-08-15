r"""
project_sup/helping_code/compare_tool.py — Unified MongoDB Comparison & Synchronization Suite.

Merges and extends:
1. Status Change Analysis (Gambling -> Regular/Blocked/Dead transitions)
2. Full Status Transition Matrix between two databases
3. Missing Domain Comparison & Sync (domain_Listed DB1 <-> DB2)
4. Missing Classification Results (checked_domains DB1 <-> DB2)
5. Single-Database Internal Audit (domain_Listed <-> checked_domains sync)

All outputs are exported as clean, clickable Excel workbooks.
Supports both interactive terminal menus and CLI flags.
"""
import os
import sys
import argparse
from pathlib import Path
from datetime import datetime, timezone, timedelta
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from dotenv import load_dotenv
from pymongo import MongoClient

# Load .env
ROOT_DIR = Path(__file__).resolve().parent.parent.parent
load_dotenv(dotenv_path=ROOT_DIR / ".env")

MONGO_URI = os.getenv("MONGO_URI", "mongodb://localhost:27017/")
DEFAULT_DB1 = os.getenv("MONGO_DB_NAME") or "gamblingsitetry"
DEFAULT_DB2 = os.getenv("MONGO_DB2_NAME") or "gamblingsites"

if DEFAULT_DB1 == DEFAULT_DB2:
    DEFAULT_DB2 = "gamblingsites" if DEFAULT_DB1 == "gamblingsitetry" else "gamblingsitetry"

COL_LISTED = os.getenv("MONGO_COLLECTION", "domain_Listed")
COL_CHECKED = os.getenv("CHECKED_COLLECTION", "checked_domains")

OUTPUT_DIR = Path(os.getenv("EXCEL_OUTPUT_DIR", "output"))
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

HYPERLINK_FONT = Font(name="Calibri", size=11, color="0000FF", underline="single")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
THIN_BORDER = Border(
    left=Side(style="thin", color="D3D3D3"),
    right=Side(style="thin", color="D3D3D3"),
    top=Side(style="thin", color="D3D3D3"),
    bottom=Side(style="thin", color="D3D3D3"),
)


def make_excel_urls_clickable(workbook_path: str, domain_column_names=("domain", "url", "domain name")):
    """Format domain/URL columns in all sheets of an Excel file to be clickable hyperlinks."""
    try:
        wb = openpyxl.load_workbook(workbook_path)
        for sheetname in wb.sheetnames:
            ws = wb[sheetname]
            target_col_indices = []

            for col_idx, cell in enumerate(ws[1], 1):
                col_name = str(cell.value or "").strip().lower()
                cell.font = HEADER_FONT
                cell.fill = HEADER_FILL
                cell.alignment = Alignment(horizontal="center", vertical="center")
                if col_name in domain_column_names:
                    target_col_indices.append(col_idx)

            for row in range(2, ws.max_row + 1):
                for col_idx in range(1, ws.max_column + 1):
                    cell = ws.cell(row=row, column=col_idx)
                    cell.border = THIN_BORDER
                    if col_idx in target_col_indices:
                        val = str(cell.value or "").strip()
                        if val and val.lower() != "none" and val.lower() != "nan":
                            target = val if val.startswith("http://") or val.startswith("https://") else f"https://{val}"
                            cell.hyperlink = target
                            cell.font = HYPERLINK_FONT

            for col in ws.columns:
                max_len = max(len(str(cell.value or "")) for cell in col)
                col_letter = openpyxl.utils.get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 60)

        wb.save(workbook_path)
    except Exception as e:
        print(f"[!] Warning: Could not apply styling/hyperlinks to {workbook_path}: {e}")


def _ask_yes_no(prompt: str) -> bool:
    while True:
        ans = input(f"{prompt} [y/n]: ").strip().lower()
        if ans in ("y", "yes"):
            return True
        if ans in ("n", "no"):
            return False
        print("    Please enter 'y' or 'n'.")


def get_mongo_client(uri: str = MONGO_URI) -> MongoClient:
    return MongoClient(uri, serverSelectionTimeoutMS=4000)


# ─────────────────────────────────────────────────────────────────────────────
# OPTION 1: Gambling -> Regular / Blocked / Dead Transitions
# ─────────────────────────────────────────────────────────────────────────────
def compare_gambling_to_other_statuses(db1_name: str, db2_name: str, client: MongoClient) -> Path | None:
    print("\n" + "=" * 70)
    print(f" [OPTION 1] COMPARING GAMBLING -> OTHER STATUS TRANSITIONS")
    print(f"  Current DB  (DB1): {db1_name}")
    print(f"  Previous DB (DB2): {db2_name}")
    print("=" * 70)

    col1 = client[db1_name][COL_CHECKED]
    col2 = client[db2_name][COL_CHECKED]

    print(f"[+] Loading {db1_name}.{COL_CHECKED}...")
    current_docs = {d["_id"]: d for d in col1.find({}, {"_id": 1, "status": 1, "reason": 1})}
    print(f"    -> Loaded {len(current_docs):,} current checked records.")

    print(f"[+] Loading {db2_name}.{COL_CHECKED}...")
    before_docs = {d["_id"]: d for d in col2.find({"status": "gambling"}, {"_id": 1, "status": 1, "reason": 1})}
    print(f"    -> Loaded {len(before_docs):,} previously gambling records.")

    changes = []
    for domain, doc_before in before_docs.items():
        doc_current = current_docs.get(domain)
        if doc_current:
            cur_status = doc_current.get("status")
            if cur_status in ("regular", "blocked", "dead"):
                raw_reason = doc_current.get("reason", [])
                reason_str = ", ".join(raw_reason) if isinstance(raw_reason, list) else str(raw_reason)
                changes.append({
                    "Domain": domain,
                    "Previous Status": "gambling",
                    "Current Status": cur_status,
                    "Current Reason": reason_str,
                })

    print(f"\n[+] Found {len(changes):,} domains that changed from 'gambling' -> 'regular'/'blocked'/'dead'.")
    if not changes:
        print("[i] No status changes found.")
        return None

    df = pd.DataFrame(changes)
    df.insert(0, "S.No.", range(1, len(df) + 1))
    out_file = OUTPUT_DIR / "status_changes_compare.xlsx"

    with pd.ExcelWriter(out_file, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="Status Changes", index=False)

    make_excel_urls_clickable(str(out_file))

    counts = df["Current Status"].value_counts().to_dict()
    print("\n" + "-" * 50)
    print(f" Summary of Status Changes:")
    print(f"  * Changed to Regular : {counts.get('regular', 0):,}")
    print(f"  * Changed to Blocked : {counts.get('blocked', 0):,}")
    print(f"  * Changed to Dead    : {counts.get('dead', 0):,}")
    print("-" * 50)
    print(f"[OK] Report exported: {out_file.resolve()}")
    return out_file


# ─────────────────────────────────────────────────────────────────────────────
# OPTION 2: Full Cross-Database Status Transition Matrix
# ─────────────────────────────────────────────────────────────────────────────
def compare_full_status_matrix(db1_name: str, db2_name: str, client: MongoClient) -> Path | None:
    print("\n" + "=" * 70)
    print(f" [OPTION 2] FULL CROSS-DATABASE STATUS TRANSITION MATRIX")
    print(f"  DB1: {db1_name}  <-->  DB2: {db2_name}")
    print("=" * 70)

    col1 = client[db1_name][COL_CHECKED]
    col2 = client[db2_name][COL_CHECKED]

    docs1 = {d["_id"]: d for d in col1.find({}, {"_id": 1, "status": 1, "reason": 1})}
    docs2 = {d["_id"]: d for d in col2.find({}, {"_id": 1, "status": 1, "reason": 1})}

    print(f"[+] Records in DB1 ({db1_name}): {len(docs1):,}")
    print(f"[+] Records in DB2 ({db2_name}): {len(docs2):,}")

    gambling_to_other = []
    other_to_gambling = []
    other_status_changes = []

    common_ids = set(docs1.keys()) & set(docs2.keys())
    print(f"[+] Common domains checked in both DBs: {len(common_ids):,}")

    for dom in common_ids:
        st1 = docs1[dom].get("status", "unknown")
        st2 = docs2[dom].get("status", "unknown")

        if st1 != st2:
            r1 = docs1[dom].get("reason", [])
            r2 = docs2[dom].get("reason", [])
            r1_str = ", ".join(r1) if isinstance(r1, list) else str(r1)
            r2_str = ", ".join(r2) if isinstance(r2, list) else str(r2)

            row = {
                "Domain": dom,
                f"Status ({db1_name})": st1,
                f"Status ({db2_name})": st2,
                f"Reason ({db1_name})": r1_str,
                f"Reason ({db2_name})": r2_str,
            }

            if st2 == "gambling" and st1 != "gambling":
                gambling_to_other.append(row)
            elif st1 == "gambling" and st2 != "gambling":
                other_to_gambling.append(row)
            else:
                other_status_changes.append(row)

    out_file = OUTPUT_DIR / "db_status_matrix_comparison.xlsx"
    with pd.ExcelWriter(out_file, engine="openpyxl") as writer:
        if gambling_to_other:
            df1 = pd.DataFrame(gambling_to_other)
            df1.insert(0, "S.No.", range(1, len(df1) + 1))
            df1.to_excel(writer, sheet_name="Gambling_to_Other", index=False)

        if other_to_gambling:
            df2 = pd.DataFrame(other_to_gambling)
            df2.insert(0, "S.No.", range(1, len(df2) + 1))
            df2.to_excel(writer, sheet_name="Other_to_Gambling", index=False)

        if other_status_changes:
            df3 = pd.DataFrame(other_status_changes)
            df3.insert(0, "S.No.", range(1, len(df3) + 1))
            df3.to_excel(writer, sheet_name="Other_Transitions", index=False)

        summary_data = [
            {"Category": f"Total records in {db1_name}", "Count": len(docs1)},
            {"Category": f"Total records in {db2_name}", "Count": len(docs2)},
            {"Category": "Common checked domains", "Count": len(common_ids)},
            {"Category": "Domains: Gambling -> Non-Gambling", "Count": len(gambling_to_other)},
            {"Category": "Domains: Non-Gambling -> Gambling", "Count": len(other_to_gambling)},
            {"Category": "Domains: Other status shifts (Dead/Blocked/Regular)", "Count": len(other_status_changes)},
            {"Category": "Total Discrepant Domains", "Count": len(gambling_to_other) + len(other_to_gambling) + len(other_status_changes)},
        ]
        pd.DataFrame(summary_data).to_excel(writer, sheet_name="Summary", index=False)

    make_excel_urls_clickable(str(out_file))

    print("\n" + "=" * 50)
    print("         TRANSITION MATRIX SUMMARY")
    print("=" * 50)
    for s in summary_data:
        print(f"  * {s['Category']:<45}: {s['Count']:,}")
    print("=" * 50)
    print(f"[OK] Full Matrix Report saved: {out_file.resolve()}")
    return out_file


# ─────────────────────────────────────────────────────────────────────────────
# OPTION 3: Compare & Synchronize domain_Listed between DB1 & DB2
# ─────────────────────────────────────────────────────────────────────────────
def compare_and_sync_domain_listed(db1_name: str, db2_name: str, client: MongoClient) -> Path | None:
    print("\n" + "=" * 70)
    print(f" [OPTION 3] COMPARE & SYNC domain_Listed BETWEEN DATABASES")
    print(f"  DB1: {db1_name}.{COL_LISTED}")
    print(f"  DB2: {db2_name}.{COL_LISTED}")
    print("=" * 70)

    col1 = client[db1_name][COL_LISTED]
    col2 = client[db2_name][COL_LISTED]

    print(f"[+] Reading domain IDs from {db1_name}.{COL_LISTED}...")
    set1 = set(d["_id"] for d in col1.find({}, {"_id": 1}))
    print(f"    -> DB1 Unique Domains: {len(set1):,}")

    print(f"[+] Reading domain IDs from {db2_name}.{COL_LISTED}...")
    set2 = set(d["_id"] for d in col2.find({}, {"_id": 1}))
    print(f"    -> DB2 Unique Domains: {len(set2):,}")

    only_in_db2 = set2 - set1
    only_in_db1 = set1 - set2

    print(f"\n[=] Domains in '{db2_name}' NOT present in '{db1_name}': {len(only_in_db2):,}")
    print(f"[=] Domains in '{db1_name}' NOT present in '{db2_name}': {len(only_in_db1):,}")

    out_file = OUTPUT_DIR / "missing_domains_compare.xlsx"
    with pd.ExcelWriter(out_file, engine="openpyxl") as writer:
        if only_in_db2:
            df2 = pd.DataFrame([{"Domain": d, "Missing In": db1_name, "Source DB": db2_name} for d in sorted(only_in_db2)])
            df2.insert(0, "S.No.", range(1, len(df2) + 1))
            df2.to_excel(writer, sheet_name=f"In_{db2_name[:15]}_Not_In_DB1", index=False)

        if only_in_db1:
            df1 = pd.DataFrame([{"Domain": d, "Missing In": db2_name, "Source DB": db1_name} for d in sorted(only_in_db1)])
            df1.insert(0, "S.No.", range(1, len(df1) + 1))
            df1.to_excel(writer, sheet_name=f"In_{db1_name[:15]}_Not_In_DB2", index=False)

        summary_df = pd.DataFrame([
            {"Metric": f"Total domains in {db1_name}", "Value": len(set1)},
            {"Metric": f"Total domains in {db2_name}", "Value": len(set2)},
            {"Metric": "Common domains in both DBs", "Value": len(set1 & set2)},
            {"Metric": f"Unique to {db2_name}", "Value": len(only_in_db2)},
            {"Metric": f"Unique to {db1_name}", "Value": len(only_in_db1)},
        ])
        summary_df.to_excel(writer, sheet_name="Summary", index=False)

    make_excel_urls_clickable(str(out_file))
    print(f"[OK] Comparison Excel saved: {out_file.resolve()}")

    # Interactive Sync Options
    if only_in_db2:
        if _ask_yes_no(f"\n[?] Do you want to copy {len(only_in_db2):,} missing domains from '{db2_name}' into '{db1_name}.{COL_LISTED}'?"):
            _bulk_copy_domains(col2, col1, only_in_db2, f"{db2_name}.{COL_LISTED}", f"{db1_name}.{COL_LISTED}")

    if only_in_db1:
        if _ask_yes_no(f"\n[?] Do you want to copy {len(only_in_db1):,} missing domains from '{db1_name}' into '{db2_name}.{COL_LISTED}'?"):
            _bulk_copy_domains(col1, col2, only_in_db1, f"{db1_name}.{COL_LISTED}", f"{db2_name}.{COL_LISTED}")

    return out_file


def _bulk_copy_domains(src_col, target_col, id_set: set, src_name: str, target_name: str):
    id_list = list(id_set)
    chunk_size = 5000
    inserted = 0
    today_str = datetime.now(timezone(timedelta(hours=5, minutes=30))).strftime("%Y-%m-%d")

    print(f"[+] Syncing {len(id_list):,} records from {src_name} to {target_name}...")
    for i in range(0, len(id_list), chunk_size):
        chunk = id_list[i : i + chunk_size]
        docs = list(src_col.find({"_id": {"$in": chunk}}))
        insert_ops = []
        for d in docs:
            doc_to_insert = {
                "_id": d["_id"],
                "domain": d.get("domain", d["_id"]),
                "active": d.get("active", True),
                "processed": False,
                "added_date": today_str,
            }
            insert_ops.append(doc_to_insert)
        if insert_ops:
            try:
                res = target_col.insert_many(insert_ops, ordered=False)
                inserted += len(res.inserted_ids)
            except Exception:
                pass
    print(f"[OK] Sync Complete. Successfully inserted {inserted:,} records into {target_name}.")


# ─────────────────────────────────────────────────────────────────────────────
# OPTION 4: Compare checked_domains between DB1 & DB2
# ─────────────────────────────────────────────────────────────────────────────
def compare_checked_domains(db1_name: str, db2_name: str, client: MongoClient) -> Path | None:
    print("\n" + "=" * 70)
    print(f" [OPTION 4] COMPARE checked_domains RECORDS BETWEEN DATABASES")
    print(f"  DB1: {db1_name}.{COL_CHECKED}")
    print(f"  DB2: {db2_name}.{COL_CHECKED}")
    print("=" * 70)

    col1 = client[db1_name][COL_CHECKED]
    col2 = client[db2_name][COL_CHECKED]

    set1 = set(d["_id"] for d in col1.find({}, {"_id": 1}))
    set2 = set(d["_id"] for d in col2.find({}, {"_id": 1}))

    print(f"[+] Checked domains in {db1_name}: {len(set1):,}")
    print(f"[+] Checked domains in {db2_name}: {len(set2):,}")

    missing_in_db1 = set2 - set1
    missing_in_db2 = set1 - set2

    print(f"\n[=] Checked in '{db2_name}' but missing in '{db1_name}': {len(missing_in_db1):,}")
    print(f"[=] Checked in '{db1_name}' but missing in '{db2_name}': {len(missing_in_db2):,}")

    out_file = OUTPUT_DIR / "missing_checked_domains.xlsx"
    with pd.ExcelWriter(out_file, engine="openpyxl") as writer:
        if missing_in_db1:
            df1 = pd.DataFrame([{"Domain": d, "Missing In": db1_name, "Source DB": db2_name} for d in sorted(missing_in_db1)])
            df1.insert(0, "S.No.", range(1, len(df1) + 1))
            df1.to_excel(writer, sheet_name=f"Missing_in_{db1_name[:15]}", index=False)

        if missing_in_db2:
            df2 = pd.DataFrame([{"Domain": d, "Missing In": db2_name, "Source DB": db1_name} for d in sorted(missing_in_db2)])
            df2.insert(0, "S.No.", range(1, len(df2) + 1))
            df2.to_excel(writer, sheet_name=f"Missing_in_{db2_name[:15]}", index=False)

        summary_df = pd.DataFrame([
            {"Metric": f"Checked in {db1_name}", "Value": len(set1)},
            {"Metric": f"Checked in {db2_name}", "Value": len(set2)},
            {"Metric": "Common checked domains in both", "Value": len(set1 & set2)},
            {"Metric": f"Checked in {db2_name} only", "Value": len(missing_in_db1)},
            {"Metric": f"Checked in {db1_name} only", "Value": len(missing_in_db2)},
        ])
        summary_df.to_excel(writer, sheet_name="Summary", index=False)

    make_excel_urls_clickable(str(out_file))
    print(f"[OK] Report exported to: {out_file.resolve()}")
    return out_file


# ─────────────────────────────────────────────────────────────────────────────
# OPTION 5: Single-DB Internal Audit (domain_Listed <-> checked_domains)
# ─────────────────────────────────────────────────────────────────────────────
def audit_single_database(db_name: str, client: MongoClient) -> Path | None:
    print("\n" + "=" * 70)
    print(f" [OPTION 5] SINGLE DATABASE INTERNAL AUDIT & SYNC")
    print(f"  Database: {db_name}")
    print(f"  Checking: {COL_LISTED} <--> {COL_CHECKED}")
    print("=" * 70)

    db = client[db_name]
    col_listed = db[COL_LISTED]
    col_checked = db[COL_CHECKED]

    total_listed = col_listed.count_documents({})
    total_checked = col_checked.count_documents({})
    unprocessed_listed = col_listed.count_documents({"processed": {"$ne": True}})

    print(f"[+] Total in '{COL_LISTED}': {total_listed:,}")
    print(f"[+] Total in '{COL_CHECKED}': {total_checked:,}")
    print(f"[+] Unprocessed in '{COL_LISTED}': {unprocessed_listed:,}")

    checked_ids = set(d["_id"] for d in col_checked.find({}, {"_id": 1}))
    listed_ids = set(d["_id"] for d in col_listed.find({}, {"_id": 1}))

    # 1. Domains in listed but missing in checked
    pending_check_ids = listed_ids - checked_ids

    # 2. Domains in checked but missing in listed
    orphan_checked_ids = checked_ids - listed_ids

    # 3. Domains in both, but processed is False/missing in listed
    mismatched_processed = list(col_listed.find({"_id": {"$in": list(checked_ids)}, "processed": {"$ne": True}}, {"_id": 1}))
    mismatched_ids = [d["_id"] for d in mismatched_processed]

    print(f"\n[=] Internal Audit Results for '{db_name}':")
    print(f"  * Listed domains pending classification : {len(pending_check_ids):,}")
    print(f"  * Checked domains missing from listed   : {len(orphan_checked_ids):,}")
    print(f"  * Domains with 'processed' flag mismatch : {len(mismatched_ids):,}")

    out_file = OUTPUT_DIR / f"internal_audit_{db_name}.xlsx"
    with pd.ExcelWriter(out_file, engine="openpyxl") as writer:
        if pending_check_ids:
            df1 = pd.DataFrame([{"Domain": d, "Status": "Pending Classification"} for d in sorted(pending_check_ids)])
            df1.insert(0, "S.No.", range(1, len(df1) + 1))
            df1.to_excel(writer, sheet_name="Pending_Classification", index=False)

        if orphan_checked_ids:
            df2 = pd.DataFrame([{"Domain": d, "Status": "Checked but Not Listed"} for d in sorted(orphan_checked_ids)])
            df2.insert(0, "S.No.", range(1, len(df2) + 1))
            df2.to_excel(writer, sheet_name="Orphan_Checked", index=False)

        if mismatched_ids:
            df3 = pd.DataFrame([{"Domain": d, "Current Flag": "processed != True", "Action Needed": "Update processed=True"} for d in sorted(mismatched_ids)])
            df3.insert(0, "S.No.", range(1, len(df3) + 1))
            df3.to_excel(writer, sheet_name="Processed_Flag_Mismatches", index=False)

        summary_df = pd.DataFrame([
            {"Metric": "Total domain_Listed", "Value": total_listed},
            {"Metric": "Total checked_domains", "Value": total_checked},
            {"Metric": "Listed domains pending checking", "Value": len(pending_check_ids)},
            {"Metric": "Checked domains orphan", "Value": len(orphan_checked_ids)},
            {"Metric": "Processed flag mismatches needing sync", "Value": len(mismatched_ids)},
        ])
        summary_df.to_excel(writer, sheet_name="Summary", index=False)

    make_excel_urls_clickable(str(out_file))
    print(f"[OK] Internal Audit Excel report saved: {out_file.resolve()}")

    if mismatched_ids:
        if _ask_yes_no(f"\n[?] Do you want to sync 'processed: True' for all {len(mismatched_ids):,} verified domains in '{COL_LISTED}'?"):
            chunk_size = 5000
            updated = 0
            for i in range(0, len(mismatched_ids), chunk_size):
                chunk = mismatched_ids[i : i + chunk_size]
                res = col_listed.update_many({"_id": {"$in": chunk}}, {"$set": {"processed": True}})
                updated += res.modified_count
            print(f"[OK] Successfully updated {updated:,} domains in '{COL_LISTED}' to processed=True.")

    return out_file


# ─────────────────────────────────────────────────────────────────────────────
# OPTION 6: Run Full Comparison Suite
# ─────────────────────────────────────────────────────────────────────────────
def run_all_comparisons(db1_name: str, db2_name: str, client: MongoClient):
    print("\n" + "=" * 70)
    print("           RUNNING FULL COMPARISON & DIAGNOSTIC SUITE")
    print("=" * 70)
    f1 = compare_gambling_to_other_statuses(db1_name, db2_name, client)
    f2 = compare_full_status_matrix(db1_name, db2_name, client)
    f3 = compare_and_sync_domain_listed(db1_name, db2_name, client)
    f4 = compare_checked_domains(db1_name, db2_name, client)
    f5 = audit_single_database(db1_name, client)
    print("\n" + "=" * 70)
    print(" [OK] FULL COMPARISON SUITE COMPLETE!")
    print(f" All reports generated in directory: {OUTPUT_DIR.resolve()}")
    print("=" * 70)


# ─────────────────────────────────────────────────────────────────────────────
# Interactive Menu CLI
# ─────────────────────────────────────────────────────────────────────────────
def interactive_menu():
    client = get_mongo_client(MONGO_URI)
    db1 = DEFAULT_DB1
    db2 = DEFAULT_DB2

    while True:
        print("\n" + "=" * 70)
        print("          MONGODB ADVANCED COMPARISON & SYNC SUITE")
        print("=" * 70)
        print(f"  Target Databases: [DB1: {db1}]  <--->  [DB2: {db2}]")
        print("=" * 70)
        print("  1. Status Shifts: Gambling -> Regular / Blocked / Dead (Clickable Excel)")
        print("  2. Full Status Matrix: Cross-Database Differences (All Status Shifts)")
        print("  3. Domain Discrepancies: Compare domain_Listed (Missing Domains & Sync)")
        print("  4. Checked Records Discrepancies: Compare checked_domains between DBs")
        print("  5. Internal Collection Sync Audit: domain_Listed vs checked_domains")
        print("  6. Run ALL Comparisons & Export Complete Excel Suite")
        print("  7. Change Target Databases (Current: DB1=" + db1 + ", DB2=" + db2 + ")")
        print("  0. Exit")
        print("=" * 70)

        choice = input("Select an option (0-7) [default: 1]: ").strip()
        if not choice:
            choice = "1"

        if choice == "1":
            compare_gambling_to_other_statuses(db1, db2, client)
        elif choice == "2":
            compare_full_status_matrix(db1, db2, client)
        elif choice == "3":
            compare_and_sync_domain_listed(db1, db2, client)
        elif choice == "4":
            compare_checked_domains(db1, db2, client)
        elif choice == "5":
            audit_single_database(db1, client)
        elif choice == "6":
            run_all_comparisons(db1, db2, client)
        elif choice == "7":
            new_db1 = input(f"Enter DB1 name [current: {db1}]: ").strip()
            new_db2 = input(f"Enter DB2 name [current: {db2}]: ").strip()
            if new_db1:
                db1 = new_db1
            if new_db2:
                db2 = new_db2
            print(f"[i] Updated target databases: DB1={db1}, DB2={db2}")
        elif choice in ("0", "exit", "quit", "q"):
            print("\nExiting comparison suite. Goodbye!")
            break
        else:
            print("[!] Invalid option. Please select 0-7.")


def main():
    parser = argparse.ArgumentParser(description="MongoDB Advanced Comparison and Sync Suite")
    parser.add_argument("--mode", "-m", choices=["1", "2", "3", "4", "5", "6", "all"], help="Comparison mode to run non-interactively")
    parser.add_argument("--db1", default=DEFAULT_DB1, help=f"Primary Database name (default: {DEFAULT_DB1})")
    parser.add_argument("--db2", default=DEFAULT_DB2, help=f"Secondary Database name (default: {DEFAULT_DB2})")
    parser.add_argument("--uri", default=MONGO_URI, help="MongoDB connection URI")

    args = parser.parse_args()

    client = get_mongo_client(args.uri)
    if args.mode:
        if args.mode == "1":
            compare_gambling_to_other_statuses(args.db1, args.db2, client)
        elif args.mode == "2":
            compare_full_status_matrix(args.db1, args.db2, client)
        elif args.mode == "3":
            compare_and_sync_domain_listed(args.db1, args.db2, client)
        elif args.mode == "4":
            compare_checked_domains(args.db1, args.db2, client)
        elif args.mode == "5":
            audit_single_database(args.db1, client)
        elif args.mode in ("6", "all"):
            run_all_comparisons(args.db1, args.db2, client)
    else:
        interactive_menu()


if __name__ == "__main__":
    main()
