"""
export_domains/exporter.py — High-Performance Direct Batch Export Pipeline
===========================================================================

Generates:
  1. Size-bounded PDF batches directly via PyMuPDF (<=24 MB, 2 targets/page, clickable URLs).
  2. Excel workbooks for each batch (clickable Domain and URL hyperlinks).
  3. Master report.xlsx (Captured Domains + Failed Domains).

Pipeline flow:
  1. Query MongoDB for unexported gambling domains.
  2. Verify screenshots via fast O(1) in-memory index (<1 sec).
  3. Move verified screenshots -> output/<run_id>/screenshots/.
  4. Generate size-bounded Batch PDFs & Excel workbooks directly in output/<run_id>/Batches/.
  5. Generate Master report.xlsx.
  6. Mark exported=True in MongoDB.
"""
from __future__ import annotations

import os
import sys
import time
import shutil
import asyncio
import urllib.parse
from datetime import datetime, timezone, timedelta
from pathlib import Path

# Ensure project root is in sys.path
_ROOT = str(Path(__file__).resolve().parent.parent)
if _ROOT not in sys.path:
    sys.path.insert(0, _ROOT)

import pymupdf
import pandas as pd
import openpyxl
from openpyxl.styles import Font
from dotenv import load_dotenv

load_dotenv(dotenv_path=Path(__file__).resolve().parent.parent / ".env")

from export_domains.screenshot import _url_to_filename, all_filename_candidates, find_screenshot_path
from db.mongo_client import checked_domains, source_domains

IST = timezone(timedelta(hours=5, minutes=30))
HYPERLINK_FONT = Font(color="0000FF", underline="single")
DEFAULT_BATCH_SIZE = int(os.getenv("EXPORT_BATCH_SIZE", 200))
DEFAULT_MAX_PDF_MB = float(os.getenv("EXPORT_MAX_PDF_MB", 23.5))


def get_screenshot_dir() -> str:
    return os.getenv("SCREENSHOT_DIR", os.path.join("output", "screenshots"))


def get_output_dir() -> str:
    return os.getenv("EXPORT_OUTPUT_DIR", "output")


def partition_entries(
    entries: list[dict],
    max_mb: float = DEFAULT_MAX_PDF_MB,
    max_targets: int = DEFAULT_BATCH_SIZE,
) -> list[list[dict]]:
    """
    Partition entries into batches ensuring each batch PDF stays strictly under max_mb (default 23.5 MB).
    Each entry adds its screenshot file size + ~2.5KB PDF page overhead.
    """
    max_bytes = int(max_mb * 1024 * 1024)
    batches = []
    current_batch = []
    current_bytes = 0

    for entry in entries:
        img_path = entry.get("screenshot_path")
        img_bytes = 0
        if img_path and os.path.exists(img_path):
            try:
                img_bytes = os.path.getsize(img_path)
            except OSError:
                img_bytes = 115_000
        else:
            img_bytes = 115_000

        est_entry_bytes = img_bytes + 2500

        if current_batch and ((current_bytes + est_entry_bytes > max_bytes) or (len(current_batch) >= max_targets)):
            batches.append(current_batch)
            current_batch = []
            current_bytes = 0

        current_batch.append(entry)
        current_bytes += est_entry_bytes

    if current_batch:
        batches.append(current_batch)

    return batches


# ── Excel Workbook Utilities ──────────────────────────────────────────────────

def _apply_excel_hyperlinks(ws):
    """Apply blue underlined hyperlinks to Domain and URL columns."""
    url_col = domain_col = None
    for idx, cell in enumerate(ws[1], 1):
        col = str(cell.value or "").strip().lower()
        if col == "url":
            url_col = idx
        elif col == "domain":
            domain_col = idx

    for col_idx in filter(None, [url_col, domain_col]):
        for row in range(2, ws.max_row + 1):
            cell = ws.cell(row=row, column=col_idx)
            val = str(cell.value or "").strip()
            if val and not val.startswith("#"):
                target = val if val.startswith(("http://", "https://")) else f"https://{val}"
                try:
                    if urllib.parse.urlparse(target).netloc:
                        cell.hyperlink = target
                        cell.font = HYPERLINK_FONT
                except Exception:
                    pass


def build_batch_workbook(entries: list[dict], output_path: str, start_no: int = 1) -> str:
    """Build batch Excel workbook with clickable Domain and URL columns."""
    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    rows = [
        {
            "S.No.": start_no + i,
            "Domain": e.get("domain", ""),
            "URL": e.get("url", ""),
        }
        for i, e in enumerate(entries)
    ]
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        pd.DataFrame(rows).to_excel(writer, sheet_name="Captured Domains", index=False)

    wb = openpyxl.load_workbook(output_path)
    if "Captured Domains" in wb.sheetnames:
        _apply_excel_hyperlinks(wb["Captured Domains"])
    wb.save(output_path)
    wb.close()
    return output_path


def build_workbook(entries: list[dict], failed_docs: list[dict], output_path: str) -> str:
    """Build master report.xlsx with 'Captured Domains' and 'Failed Domains' sheets."""
    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)

    captured_rows = [
        {
            "S.No.": i,
            "Domain": e.get("domain", ""),
            "URL": e.get("url", ""),
        }
        for i, e in enumerate(entries, 1)
    ]

    failed_rows = [
        {
            "S.No.": i,
            "Domain": d.get("domain") or d.get("_id", ""),
            "URL": d.get("url", ""),
            "Reason": d.get("screenshot_failed_reason") or d.get("reason", "Screenshot missing"),
        }
        for i, d in enumerate(failed_docs, 1)
    ]

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        pd.DataFrame(captured_rows).to_excel(writer, sheet_name="Captured Domains", index=False)
        if failed_rows:
            pd.DataFrame(failed_rows).to_excel(writer, sheet_name="Failed Domains", index=False)

    wb = openpyxl.load_workbook(output_path)
    for sheet_name in wb.sheetnames:
        _apply_excel_hyperlinks(wb[sheet_name])
    wb.save(output_path)
    wb.close()
    return output_path


# ── PyMuPDF Fast PDF Generation ───────────────────────────────────────────────

def _draw_target(page, entry: dict, target_num: int, is_top: bool = True):
    """Draw a single target (header, clickable link, date, centered image) on the page half."""
    y_base = 40 if is_top else 436

    # 1. Target Header
    page.insert_text(
        (36, y_base + 4),
        f"TARGET #{target_num:05d}",
        fontname="helv",
        fontsize=11,
        color=(0.12, 0.16, 0.22),
    )

    # 2. Clickable URL
    url = entry.get("url") or ""
    full_url = url if (url.startswith("http://") or url.startswith("https://")) else f"https://{url}"
    page.insert_text((180, y_base + 4), "URL: ", fontname="helv", fontsize=11, color=(0.4, 0.4, 0.4))

    disp_url = url if len(url) <= 52 else url[:49] + "..."
    text_width = len(disp_url) * 6.2
    url_rect = pymupdf.Rect(215, y_base - 8, min(559, 215 + text_width), y_base + 8)
    page.insert_text((215, y_base + 4), disp_url, fontname="helv", fontsize=11, color=(0.0, 0.2, 0.8))
    page.draw_line((215, y_base + 6), (min(559, 215 + text_width), y_base + 6), color=(0.0, 0.2, 0.8), width=0.6)
    try:
        page.insert_link({"kind": pymupdf.LINK_URI, "from": url_rect, "uri": full_url})
    except Exception:
        pass

    # 3. Screenshot Date
    ss_date_raw = str(entry.get("screenshot_date") or entry.get("added_date") or datetime.now(IST).strftime("%Y-%m-%d"))
    ss_date = ss_date_raw.split(" ")[0].split("T")[0]
    page.insert_text((36, y_base + 20), "SCREENSHOT DATE: ", fontname="helv", fontsize=8.5, color=(0.45, 0.45, 0.45))
    page.insert_text((130, y_base + 20), ss_date, fontname="helv", fontsize=8.5, color=(0.2, 0.2, 0.2))

    # 4. Centered Screenshot Image
    img_path = entry.get("screenshot_path")
    y_img_top = y_base + 26
    y_img_bot = y_base + 355
    img_rect = pymupdf.Rect(36, y_img_top, 559, y_img_bot)

    if img_path and os.path.exists(img_path):
        try:
            page.insert_image(img_rect, filename=img_path, keep_proportion=True)
        except Exception as e:
            page.insert_text(
                (40, y_img_top + 40),
                f"[Image Render Error: {e}]",
                fontname="helv",
                fontsize=9,
                color=(0.8, 0.1, 0.1),
            )


def generate_pdf_batch(entries: list[dict], output_pdf_path: str, start_target_num: int = 1) -> str:
    """Generate a high-resolution, compressed A4 PDF report for an entry batch using PyMuPDF."""
    os.makedirs(os.path.dirname(os.path.abspath(output_pdf_path)), exist_ok=True)
    doc = pymupdf.open()
    total = len(entries)

    for i in range(0, total, 2):
        page = doc.new_page(width=595.44, height=841.68)  # Standard A4 in points
        # Top target
        _draw_target(page, entries[i], start_target_num + i, is_top=True)
        # Bottom target
        if i + 1 < total:
            # Subtle divider line between top and bottom targets
            page.draw_line((36, 416), (559, 416), color=(0.88, 0.90, 0.92), width=0.5)
            _draw_target(page, entries[i + 1], start_target_num + i + 1, is_top=False)

    doc.save(output_pdf_path, garbage=4, deflate=True)
    doc.close()
    return output_pdf_path


# ── Unified Export Pipeline ───────────────────────────────────────────────────

def run_export(
    domain_ids: list = None,
    limit: int = 0,
    batch_size: int = DEFAULT_BATCH_SIZE,
    max_mb: float = DEFAULT_MAX_PDF_MB,
) -> dict:
    """Single-command high-speed export: Direct size-bounded PDF & Excel batches."""
    t_start = time.time()

    # 1. Query MongoDB
    query = {"status": "gambling", "exported": {"$ne": True}}
    if domain_ids:
        query["_id"] = {"$in": domain_ids}

    docs = list(checked_domains().find(query))
    if limit:
        docs = docs[:limit]

    if not docs:
        print("[export] No unexported gambling domains found — nothing to do.")
        return {"run_id": None, "pdf": None, "xlsx": None, "captured": 0, "failed": 0, "batches": 0}

    # 2. Fast O(1) indexing of disk files
    src_dir = get_screenshot_dir()
    disk_files: dict[str, str] = {}
    if os.path.isdir(src_dir):
        for fname in os.listdir(src_dir):
            fl = fname.lower()
            if fl.endswith((".jpg", ".png", ".jpeg", ".webp")):
                fp = os.path.join(src_dir, fname)
                try:
                    if os.path.getsize(fp) >= 4000:
                        disk_files[fl] = fp
                except OSError:
                    pass

    print(f"[export] Indexed {len(disk_files):,} valid screenshots on disk in {time.time() - t_start:.2f}s.")

    # 3. Setup output destination
    run_id = datetime.now(IST).strftime("%Y%m%d_%H%M%S")
    dest_dir = os.path.join(get_output_dir(), run_id)
    dest_screenshots = os.path.join(dest_dir, "screenshots")
    os.makedirs(dest_screenshots, exist_ok=True)

    entries = []
    failed_docs = []
    missing_ids = []

    print(f"[export] Verifying and moving screenshots for {len(docs):,} domains...")
    t_move = time.time()

    for doc in docs:
        domain = doc.get("domain") or doc.get("_id")
        url = doc.get("url") or f"https://{domain}"

        # Fast candidate lookup
        src = None
        for cand in all_filename_candidates(url, domain):
            cand_l = cand.lower()
            if cand_l in disk_files:
                src = disk_files[cand_l]
                break

        # Fallback to general find_screenshot_path if not in flat indexed set
        if not src:
            found = find_screenshot_path(url, domain, src_dir)
            if found and os.path.exists(found) and os.path.getsize(found) >= 4000:
                src = found

        if src:
            dest = os.path.join(dest_screenshots, os.path.basename(src))
            ss_date = doc.get("screenshot_date") or doc.get("added_date")
            entry_item = {
                "domain": domain,
                "url": url,
                "_id": doc["_id"],
                "_src": src,
                "screenshot_date": ss_date,
                "ip": doc.get("ip", ""),
            }
            try:
                if os.path.abspath(src) != os.path.abspath(dest):
                    if os.path.exists(dest):
                        os.remove(dest)
                    shutil.move(src, dest)
                entry_item["screenshot_path"] = dest
                entries.append(entry_item)
            except Exception as e:
                if os.path.exists(dest):
                    entry_item["screenshot_path"] = dest
                    entries.append(entry_item)
                elif os.path.exists(src):
                    print(f"  [export] Move failed for {domain}: {e} — using source path")
                    entry_item["screenshot_path"] = src
                    entries.append(entry_item)
                else:
                    failed_docs.append(doc)
        else:
            if doc.get("screenshot_taken"):
                missing_ids.append(doc["_id"])
            failed_docs.append(doc)

    if missing_ids:
        print(f"  [export] {len(missing_ids)} domain(s) had screenshot missing — marking status as 'unconfirmed'.")
        checked_domains().update_many(
            {"_id": {"$in": missing_ids}},
            {"$set": {
                "status": "unconfirmed",
                "screenshot_taken": False,
                "reason": "Unconfirmed: Screenshot file missing at export time",
                "screenshot_failed_reason": "Screenshot file missing at export time",
                "ai_evaluated": False,
            }},
        )

    print(f"[export] Move & Verification complete: Verified={len(entries):,} | Failed={len(failed_docs):,} ({time.time() - t_move:.2f}s)")

    if not entries:
        print("[export] No valid screenshots found. Nothing to export.")
        return {"run_id": run_id, "pdf": None, "xlsx": None, "captured": 0, "failed": len(failed_docs), "batches": 0}

    # 4. Generate Batches Directly (<=24 MB per PDF)
    batches_dir = os.path.join(dest_dir, "Batches")
    os.makedirs(batches_dir, exist_ok=True)

    batches = partition_entries(entries, max_mb=max_mb, max_targets=batch_size)
    num_batches = len(batches)
    print(f"\n[export] Compiling {num_batches} batches directly via PyMuPDF (<= {max_mb} MB, up to {batch_size} targets/batch)...")

    batch_folders = []
    first_pdf = None
    cumulative_target_count = 0

    for b_idx, batch_slice in enumerate(batches, 1):
        folder_name = f"Batch {b_idx}"
        folder_path = os.path.join(batches_dir, folder_name)
        os.makedirs(folder_path, exist_ok=True)
        batch_folders.append(folder_path)

        start_target_num = cumulative_target_count + 1

        # Batch PDF
        b_pdf = os.path.join(folder_path, f"batch_{b_idx}.pdf")
        generate_pdf_batch(batch_slice, b_pdf, start_target_num=start_target_num)
        if not first_pdf:
            first_pdf = b_pdf

        # Batch Excel
        b_xlsx = os.path.join(folder_path, f"batch_{b_idx}.xlsx")
        build_batch_workbook(batch_slice, b_xlsx, start_no=start_target_num)

        pdf_size_mb = os.path.getsize(b_pdf) / (1024 * 1024)
        print(f"  [+] {folder_name:10} : {len(batch_slice):3} targets (#{start_target_num:05d} - #{start_target_num + len(batch_slice) - 1:05d}) | PDF: {pdf_size_mb:4.1f} MB -> {b_pdf}")
        cumulative_target_count += len(batch_slice)

    # 5. Build Master Excel Workbook
    xlsx_path = os.path.join(dest_dir, "report.xlsx")
    print(f"\n[export] Building Master Excel workbook ({len(entries):,} captured, {len(failed_docs):,} failed)...")
    build_workbook(entries, failed_docs, xlsx_path)
    print(f"[export] Master Excel saved: {xlsx_path}")

    # 6. Mark exported=True in MongoDB
    exported_ids = [e["_id"] for e in entries]
    now_ist = datetime.now(IST).strftime("%Y-%m-%d")
    checked_domains().update_many(
        {"_id": {"$in": exported_ids}},
        {"$set": {"exported": True, "exported_at": now_ist, "screenshot_taken": True}},
    )

    elapsed_total = time.time() - t_start
    print(f"\n" + "=" * 65)
    print(f"  EXPORT COMPLETED IN {elapsed_total:.1f}s (~{elapsed_total/60:.1f} min)")
    print("=" * 65)
    print(f"  Run Directory : {dest_dir}")
    print(f"  Batches       : {num_batches} batch folders in 'Batches/'")
    print(f"  Master Excel  : {xlsx_path}")
    print(f"  Screenshots   : {dest_screenshots}")
    print(f"  Captured      : {len(entries):,} | Failed: {len(failed_docs):,}")
    print("=" * 65 + "\n")

    return {
        "run_id": run_id,
        "pdf": first_pdf,
        "xlsx": xlsx_path,
        "captured": len(entries),
        "failed": len(failed_docs),
        "batches": num_batches,
        "batches_dir": batches_dir,
    }


async def run(concurrency: int = None, limit: int = 0) -> dict:
    """Async entry point."""
    return await asyncio.to_thread(run_export, None, limit)


# Backward compatibility aliases
build_report = generate_pdf_batch
_convert_to_pdf = lambda docx_path: docx_path.replace(".docx", ".pdf")


if __name__ == "__main__":
    import argparse
    parser = argparse.ArgumentParser(description="High-Speed Direct Batch Exporter")
    parser.add_argument("--limit", type=int, default=int(os.getenv("CAPTURE_LIMIT", 0)))
    parser.add_argument("--batch-size", type=int, default=DEFAULT_BATCH_SIZE, help=f"Max targets per batch (default: {DEFAULT_BATCH_SIZE})")
    parser.add_argument("--max-mb", type=float, default=DEFAULT_MAX_PDF_MB, help=f"Max PDF size in MB (default: {DEFAULT_MAX_PDF_MB})")
    args = parser.parse_args()
    run_export(limit=args.limit, batch_size=args.batch_size, max_mb=args.max_mb)
