import os
import io
from xml.sax.saxutils import escape
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml import parse_xml
from docx.opc.constants import RELATIONSHIP_TYPE

def add_clickable_hyperlink(paragraph, url: str, text: str = None, font_size_pt: float = 12.0):
    """Inject an OpenXML <w:hyperlink> into a python-docx paragraph for PDF & Word clickable links."""
    if text is None:
        text = url

    # Relate URL to document part
    r_id = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)

    # Half-points for OpenXML w:sz (12 pt -> 24 half-points)
    half_pts = int(font_size_pt * 2)

    hyperlink_xml = f'<w:hyperlink xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" r:id="{r_id}"/>'
    hyperlink = parse_xml(hyperlink_xml)

    run_xml = f'''<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
        <w:rPr>
            <w:rFonts w:ascii="Arial" w:hAnsi="Arial"/>
            <w:color w:val="0000FF"/>
            <w:u w:val="single"/>
            <w:sz w:val="{half_pts}"/>
            <w:szCs w:val="{half_pts}"/>
        </w:rPr>
        <w:t>{escape(text)}</w:t>
    </w:r>'''
    run_node = parse_xml(run_xml)
    hyperlink.append(run_node)
    paragraph._p.append(hyperlink)


class DocxReportGenerator:
    def __init__(self, output_path="data/presence_report.docx"):
        self.output_path = output_path
        os.makedirs(os.path.dirname(os.path.abspath(self.output_path)), exist_ok=True)
        self.doc = docx.Document()
        self._configure_page_setup()
        self.target_count = 0

    def _configure_page_setup(self):
        """Set page layout strictly to A4 Portrait with 0.4 in margins (2 targets per page)."""
        section = self.doc.sections[0]
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.top_margin = Inches(0.4)
        section.bottom_margin = Inches(0.4)
        section.left_margin = Inches(0.4)
        section.right_margin = Inches(0.4)

    def add_target(self, url: str, image_buffer: io.BytesIO, index: int = None):
        """Add a target item (Single-line Header + Clickable URL + 16:9 Screenshot) to the Word document."""
        self.target_count += 1
        current_idx = index if index is not None else self.target_count

        # Single Header Line: Left: TARGET #00001 | Right: URL: https://...
        p_hdr = self.doc.add_paragraph()
        p_hdr.paragraph_format.space_before = Pt(0)
        p_hdr.paragraph_format.space_after = Pt(3)
        p_hdr.paragraph_format.tab_stops.add_tab_stop(Inches(7.1), WD_TAB_ALIGNMENT.RIGHT)

        # Target ID on left (12pt Bold Arial)
        run_hdr = p_hdr.add_run(f"TARGET #{current_idx:05d}")
        run_hdr.font.name = 'Arial'
        run_hdr.font.size = Pt(12)
        run_hdr.font.bold = True
        run_hdr.font.color.rgb = RGBColor(30, 40, 55)

        # Right Tab + URL label (12pt Bold Arial)
        run_url_lbl = p_hdr.add_run("\tURL: ")
        run_url_lbl.font.name = 'Arial'
        run_url_lbl.font.size = Pt(12)
        run_url_lbl.font.bold = True
        run_url_lbl.font.color.rgb = RGBColor(60, 70, 80)

        # Clickable Hyperlink on same line (12pt Blue Underlined Arial)
        add_clickable_hyperlink(p_hdr, url=url, font_size_pt=12.0)

        # Image Paragraph
        p_img = self.doc.add_paragraph()
        p_img.paragraph_format.space_before = Pt(0)
        p_img.paragraph_format.space_after = Pt(6)
        p_img.add_run().add_picture(image_buffer, width=Inches(7.1))

        # Page Break after every 2 targets
        if self.target_count % 2 == 0:
            self.doc.add_page_break()

    def save(self):
        """Save the generated Word document to disk."""
        self.doc.save(self.output_path)
        return self.output_path
