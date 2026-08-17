import asyncio
import os
from datetime import datetime, timezone, timedelta
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from dotenv import load_dotenv

load_dotenv(dotenv_path=Path(__file__).resolve().parent.parent / ".env")

from db.mongo_client import find_pending_capture, checked_domains
from capture_url.screenshot import BrowserPool, is_valid_screenshot

OUTPUT_DIR = os.getenv("SCREENSHOT_DIR", os.path.join("output", "screenshots"))
CONCURRENCY = int(os.getenv("SCREENSHOT_CONCURRENCY", 35))
IST = timezone(timedelta(hours=5, minutes=30))

# ── Styling Constants ─────────────────────────────────────────────────────────
HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

KPI_TITLE_FILL = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
KPI_TITLE_FONT = Font(name="Calibri", size=12, bold=True, color="FFFFFF")

SUCCESS_FILL = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
SUCCESS_FONT = Font(name="Calibri", size=11, bold=True, color="375623")

FAIL_FILL = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
FAIL_FONT = Font(name="Calibri", size=11, bold=True, color="C65911")

WARN_FILL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
WARN_FONT = Font(name="Calibri", size=11, bold=True, color="806000")

HYPERLINK_FONT = Font(name="Calibri", size=11, color="0000FF", underline="single")
REGULAR_FONT = Font(name="Calibri", size=11)
BOLD_FONT = Font(name="Calibri", size=11, bold=True)

THIN_BORDER = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9"),
)


def export_screenshot_summary_excel(records: list[dict], output_excel_path: str = None) -> str:
    """Generate a clean, styled multi-tab Excel summary workbook of all processed domains."""
    now_ist = datetime.now(IST)
    run_ts = now_ist.strftime("%Y%m%d_%H%M%S")

    if not output_excel_path:
        out_folder = os.path.join("output", run_ts)
        os.makedirs(out_folder, exist_ok=True)
        output_excel_path = os.path.join(out_folder, f"screenshot_capture_summary_{run_ts}.xlsx")
    else:
        os.makedirs(os.path.dirname(output_excel_path), exist_ok=True)

    wb = openpyxl.Workbook()
    # Remove default sheet
    default_sheet = wb.active

    # Calculate statistics
    total = len(records)
    captured = [r for r in records if r["status"] == "Captured"]
    failed = [r for r in records if r["status"] == "Failed"]
    
    blocked = [r for r in failed if r.get("category") == "blocked"]
    dead = [r for r in failed if r.get("category") == "dead"]
    other_failed = [r for r in failed if r.get("category") not in ("blocked", "dead")]

    success_rate = (len(captured) / total * 100) if total > 0 else 0.0

    # ── SHEET 1: Executive Summary ────────────────────────────────────────────
    ws_summary = wb.create_sheet(title="Executive Summary")
    ws_summary.views.sheetView[0].showGridLines = True

    # Title block
    ws_summary.merge_cells("A1:D1")
    title_cell = ws_summary["A1"]
    title_cell.value = "GAMBLINGWEBFIND — SCREENSHOT CAPTURE SUMMARY"
    title_cell.fill = KPI_TITLE_FILL
    title_cell.font = KPI_TITLE_FONT
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws_summary.row_dimensions[1].height = 28

    summary_rows = [
        ("Execution Timestamp", now_ist.strftime("%Y-%m-%d %H:%M:%S IST")),
        ("Total Domains Queued", total),
        ("Successfully Captured", len(captured)),
        ("Failed Captures", len(failed)),
        ("  • Blocked (403 / Cloudflare WAF)", len(blocked)),
        ("  • Dead (No response / Timed out / 404 / Parked)", len(dead)),
        ("  • Other Errors", len(other_failed)),
        ("Capture Success Rate", f"{success_rate:.2f}%"),
        ("Screenshot Storage Folder", os.path.abspath(OUTPUT_DIR)),
    ]

    for idx, (label, val) in enumerate(summary_rows, start=3):
        cell_lbl = ws_summary.cell(row=idx, column=1, value=label)
        cell_val = ws_summary.cell(row=idx, column=2, value=val)

        cell_lbl.font = BOLD_FONT
        cell_lbl.border = THIN_BORDER
        cell_val.font = REGULAR_FONT
        cell_val.border = THIN_BORDER

        if "Success Rate" in label:
            cell_val.font = Font(name="Calibri", size=11, bold=True, color="375623" if success_rate >= 50 else "C65911")
        elif "Successfully Captured" in label:
            cell_val.fill = SUCCESS_FILL
            cell_val.font = SUCCESS_FONT
        elif "Failed Captures" in label and len(failed) > 0:
            cell_val.fill = FAIL_FILL
            cell_val.font = FAIL_FONT

    # ── SHEET 2: All Domains ──────────────────────────────────────────────────
    ws_all = wb.create_sheet(title="All Domains")
    _populate_domain_sheet(ws_all, records)

    # ── SHEET 3: Captured (Success) ───────────────────────────────────────────
    ws_cap = wb.create_sheet(title="Captured (Success)")
    _populate_domain_sheet(ws_cap, captured)

    # ── SHEET 4: Failed (Action Required) ─────────────────────────────────────
    ws_fail = wb.create_sheet(title="Failed Domains")
    _populate_domain_sheet(ws_fail, failed)

    # Remove initial blank sheet
    if default_sheet in wb.worksheets:
        wb.remove(default_sheet)

    wb.save(output_excel_path)
    return output_excel_path


def _populate_domain_sheet(ws, items: list[dict]):
    """Format and populate domain records into a sheet with styling and hyperlinks."""
    ws.views.sheetView[0].showGridLines = True

    headers = [
        "S.No.",
        "Domain",
        "Target URL",
        "Capture Status",
        "Category",
        "Details / Failure Reason",
        "Screenshot File",
        "Timestamp",
    ]

    ws.append(headers)
    ws.row_dimensions[1].height = 24

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center" if col_idx in (1, 4, 5) else "left", vertical="center")

    for i, r in enumerate(items, start=1):
        row_num = i + 1
        ws.append([
            i,
            r.get("domain", ""),
            r.get("url", f"https://{r.get('domain', '')}"),
            r.get("status", ""),
            r.get("category", ""),
            r.get("reason", ""),
            r.get("screenshot_file", "None"),
            r.get("timestamp", ""),
        ])

        ws.row_dimensions[row_num].height = 20

        # Cells formatting
        c_sno = ws.cell(row=row_num, column=1)
        c_domain = ws.cell(row=row_num, column=2)
        c_url = ws.cell(row=row_num, column=3)
        c_status = ws.cell(row=row_num, column=4)
        c_cat = ws.cell(row=row_num, column=5)
        c_reason = ws.cell(row=row_num, column=6)
        c_file = ws.cell(row=row_num, column=7)
        c_time = ws.cell(row=row_num, column=8)

        for c in (c_sno, c_domain, c_url, c_status, c_cat, c_reason, c_file, c_time):
            c.border = THIN_BORDER
            c.font = REGULAR_FONT

        c_sno.alignment = Alignment(horizontal="center", vertical="center")
        c_cat.alignment = Alignment(horizontal="center", vertical="center")
        c_time.alignment = Alignment(horizontal="center", vertical="center")

        # Clickable Hyperlink for Domain
        dom_val = r.get("domain", "")
        if dom_val:
            c_domain.hyperlink = f"https://{dom_val}"
            c_domain.font = HYPERLINK_FONT

        # Clickable Hyperlink for Target URL
        url_val = r.get("url", "")
        if url_val:
            c_url.hyperlink = url_val
            c_url.font = HYPERLINK_FONT

        # Status badge coloring
        status_val = r.get("status", "")
        if status_val == "Captured":
            c_status.fill = SUCCESS_FILL
            c_status.font = SUCCESS_FONT
            c_status.alignment = Alignment(horizontal="center", vertical="center")
        else:
            c_status.fill = FAIL_FILL
            c_status.font = FAIL_FONT
            c_status.alignment = Alignment(horizontal="center", vertical="center")

    # Auto-adjust column widths
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 11), 60)


async def run(domain_ids=None, concurrency=None, limit=0, output_excel_path=None):
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    concurrency = concurrency or CONCURRENCY

    if domain_ids is not None:
        today = datetime.now(IST).strftime("%Y-%m-%d")
        docs = []
        for d in domain_ids:
            doc = checked_domains().find_one({"_id": d})
            if not doc:
                # Auto-insert as confirmed gambling true positive
                checked_domains().update_one(
                    {"_id": d},
                    {"$set": {
                        "domain": d,
                        "url": f"https://{d}",
                        "status": "gambling",
                        "reason": "Manual true positive import",
                        "screenshot_taken": False,
                        "screenshot_failed_reason": None,
                        "source": "manual_import",
                    }, "$setOnInsert": {"added_date": today}},
                    upsert=True,
                )
                doc = checked_domains().find_one({"_id": d})
                print(f"[screenshot_runner] Auto-inserted new domain: {d}")
            docs.append(doc)
    else:
        docs = list(find_pending_capture(limit=limit))

    if not docs:
        print("[screenshot_runner] No gambling domains pending screenshot capture.")
        return {"captured": 0, "failed": 0, "total": 0, "excel_path": None}

    print(f"[screenshot_runner] {len(docs)} domains queued for screenshot capture.")

    pool = BrowserPool(concurrency=concurrency)
    await pool.start()
    sem = asyncio.Semaphore(concurrency)
    results = {"captured": 0, "failed": 0}
    records = []

    async def _capture_one(doc):
        domain = doc.get("domain") or doc.get("_id")
        url = doc.get("url") or f"https://{domain}"
        async with sem:
            try:
                filepath, status, _ = await pool.capture_url(url, OUTPUT_DIR, retries=2)
            except Exception as e:
                filepath, status = None, "error"
                print(f"[screenshot_runner] ERROR {domain}: {e}", flush=True)

        ts_now = datetime.now(IST).strftime("%Y-%m-%d %H:%M:%S")

        if filepath and is_valid_screenshot(filepath):
            checked_domains().update_one(
                {"_id": domain},
                {"$set": {"screenshot_taken": True, "screenshot_failed_reason": None}}
            )
            results["captured"] += 1
            print(f"[screenshot_runner] OK {domain}", flush=True)
            records.append({
                "domain": domain,
                "url": url,
                "status": "Captured",
                "category": "success",
                "reason": "Screenshot verified and saved",
                "screenshot_file": os.path.basename(filepath),
                "timestamp": ts_now,
            })
        else:
            reason = f"Capture failed: {status}" if status != "error" else "Capture error (browser crash or network timeout)"
            checked_domains().update_one(
                {"_id": domain},
                {"$set": {"screenshot_taken": False, "screenshot_failed_reason": reason}}
            )
            results["failed"] += 1
            print(f"[screenshot_runner] FAIL {domain} ({status})", flush=True)
            records.append({
                "domain": domain,
                "url": url,
                "status": "Failed",
                "category": status,
                "reason": reason,
                "screenshot_file": "None",
                "timestamp": ts_now,
            })

    await asyncio.gather(*[_capture_one(doc) for doc in docs])
    await pool.close()

    # Generate and export the complete summary Excel
    excel_path = export_screenshot_summary_excel(records, output_excel_path=output_excel_path)
    results["total"] = len(docs)
    results["excel_path"] = excel_path

    print(f"\n[screenshot_runner] Complete summary Excel exported:")
    print(f"    File path : {os.path.abspath(excel_path)}")
    print(f"    Captured  : {results['captured']} / {results['total']} ({(results['captured']/results['total']*100 if results['total']>0 else 0):.1f}%)")
    print(f"    Failed    : {results['failed']}")

    return results
